import logging
import pandas as pd
import fitz  # PyMuPDF
import openpyxl

logger = logging.getLogger(__name__)


def extract_text_from_pdf(file_path: str) -> str:
    text = ""
    try:
        doc = fitz.open(file_path)
        for page in doc:
            page_text = page.get_text("text")
            if page_text:
                text += page_text + "\n\n"
        doc.close()
    except Exception as e:
        logger.error("PDF extraction failed: %s", e)
        raise
    return text


def extract_text_from_excel(file_path: str) -> str:
    text = ""
    try:
        wb = openpyxl.load_workbook(file_path, read_only=True, data_only=True)
        for sheet_name in wb.sheetnames:
            ws = wb[sheet_name]
            text += f"\n\n=== Sheet: {sheet_name} ===\n"
            for row in ws.iter_rows(values_only=True):
                row_values = [str(cell) if cell is not None else "" for cell in row]
                if any(v.strip() for v in row_values):
                    text += "\t".join(row_values) + "\n"
    except Exception as e:
        logger.error("Excel extraction failed: %s", e)
        raise
    return text


_TB_COLS = {"account name", "debit", "credit"}


def try_parse_trial_balance_csv(file_path: str) -> dict | None:
    """
    Read the file once. If it looks like a TB-tool CSV, parse and return the
    extracted_data dict. Returns None if the file is not a trial balance CSV.
    Bypasses Groq — the structure is already clean and well-defined.
    """
    try:
        df = pd.read_csv(file_path, dtype=str).fillna("")
    except Exception:
        return None

    df.columns = [c.strip() for c in df.columns]
    if not _TB_COLS.issubset({c.lower() for c in df.columns}):
        return None

    # Drop TOTAL row
    df = df[df.get("Account Name", pd.Series()).str.upper().str.strip() != "TOTAL"].copy()

    def to_float(v):
        v = str(v).strip().replace(",", "")
        try:
            return float(v) if v else None
        except ValueError:
            return None

    # Group by Account Type
    type_order = ["asset", "liability", "equity", "revenue", "expense"]
    groups: dict[str, list] = {}
    for _, row in df.iterrows():
        atype = str(row.get("Account Type", "")).strip()
        if not atype or atype.lower() == "nan":
            atype = "Other"
        key = atype.lower()
        groups.setdefault(key, []).append(row)

    sections = []
    for tkey in type_order + [k for k in groups if k not in type_order]:
        if tkey not in groups:
            continue
        items = groups[tkey]
        section_name = tkey.capitalize() + "s" if not tkey.endswith("s") else tkey.capitalize()

        line_items = []
        subtotal_debit = subtotal_credit = 0.0
        for row in items:
            debit = to_float(row.get("Debit", ""))
            credit = to_float(row.get("Credit", ""))
            # Represent as signed net for analysis: debits positive, credits negative
            net = (debit or 0.0) - (credit or 0.0)
            subtotal_debit += debit or 0.0
            subtotal_credit += credit or 0.0
            line_items.append({
                "label": str(row.get("Account Name", "")).strip(),
                "account_number": str(row.get("Account Number", "")).strip(),
                "debit": debit,
                "credit": credit,
                "current_year": net,
                "prior_year": None,
                "is_subtotal": False,
                "is_total": False,
                "indent_level": 0,
            })

        subtotal_net = subtotal_debit - subtotal_credit
        sections.append({
            "name": section_name,
            "line_items": line_items,
            "subtotal": {
                "label": f"Total {section_name}",
                "debit": subtotal_debit,
                "credit": subtotal_credit,
                "current_year": subtotal_net,
                "prior_year": None,
                "is_subtotal": True,
                "is_total": False,
                "indent_level": 0,
            },
        })

    total_debit = sum(to_float(r.get("Debit", "")) or 0 for _, r in df.iterrows())
    total_credit = sum(to_float(r.get("Credit", "")) or 0 for _, r in df.iterrows())

    raw_text = df.to_string(index=False)
    return {
        "statement_type": "trial_balance",
        "period": "",
        "currency": "USD",
        "unit": "ones",
        "sections": sections,
        "total": {
            "label": "Total",
            "debit": total_debit,
            "credit": total_credit,
            "current_year": total_debit - total_credit,
            "prior_year": None,
            "is_subtotal": False,
            "is_total": True,
            "indent_level": 0,
        },
        "_raw_text": raw_text,
    }


def extract_text_from_csv(file_path: str) -> str:
    try:
        df = pd.read_csv(file_path)
        return df.to_string(index=False)
    except Exception as e:
        logger.error("CSV extraction failed: %s", e)
        raise


def extract_document_text(file_path: str, file_type: str) -> str:
    if file_type == "pdf":
        return extract_text_from_pdf(file_path)
    elif file_type in ("xlsx", "xls"):
        return extract_text_from_excel(file_path)
    elif file_type == "csv":
        return extract_text_from_csv(file_path)
    raise ValueError(f"Unsupported file type: {file_type}")


def find_text_in_pdf(file_path: str, text: str) -> dict | None:
    """
    Search for text in a PDF and return its normalized coordinates
    {page, x, y, w, h} where all values are 0-1 fractions of page dimensions.
    Returns None if text is not found.
    """
    if not text:
        return None
    try:
        doc = fitz.open(file_path)
        for page_num, page in enumerate(doc):
            pw = page.rect.width
            ph = page.rect.height
            hits = page.search_for(text, quads=False)
            if hits:
                r = hits[0]  # take first occurrence
                doc.close()
                return {
                    "page": page_num,
                    "x": round(r.x0 / pw, 4),
                    "y": round(r.y0 / ph, 4),
                    "w": round((r.x1 - r.x0) / pw, 4),
                    "h": round((r.y1 - r.y0) / ph, 4),
                }
        doc.close()
    except Exception as e:
        logger.warning("Text search failed for '%s': %s", text, e)
    return None
